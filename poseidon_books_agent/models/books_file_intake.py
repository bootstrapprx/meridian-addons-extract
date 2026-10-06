"""Bounded journal files share the source-evidence review and approval lifecycle."""
import base64
import csv
import datetime
import io
import itertools
import math
import zipfile

from odoo import api, fields, models
from odoo.exceptions import AccessError, ValidationError
from .books_intake import BooksIntake, digest


def journal_rows(filename, encoded):
    if not isinstance(filename, str) or not isinstance(encoded, str) or len(encoded) > 5600000:
        raise ValidationError("Choose a CSV/XLSX journal file up to 4 MB.")
    try:
        content = base64.b64decode(encoded, validate=True)
        if len(content) > 4 * 1024 * 1024:
            raise ValueError("File too large")
        extension = filename.lower().rsplit(".", 1)[-1]
        if extension == "csv":
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            rows = list(itertools.islice(reader, 5001))
        elif extension == "xlsx":
            import openpyxl
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Spreadsheet expands beyond the limit")
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > 20:
                    raise ValueError("Spreadsheet exceeds row or column limits")
                iterator = sheet.iter_rows(values_only=True)
                headers = list(next(iterator, ()))
                rows = [dict(zip(headers, values)) for values in itertools.islice(iterator, 5001) if any(v is not None for v in values)]
            finally:
                workbook.close()
        else:
            raise ValueError("Unsupported format")
        required = {"entry_id", "date", "account_code", "debit", "credit"}
        if not set(headers).issubset(required | {"description", "currency"}) or len(headers) != len(set(headers)) or not required.issubset(headers) or not 1 <= len(rows) <= 5000:
            raise ValueError("Invalid journal columns or row count")
        if any(None in row for row in rows):
            raise ValueError("A row has extra columns")
        for row in rows:
            for key, value in list(row.items()):
                if isinstance(value, datetime.datetime):
                    row[key] = value.date().isoformat()
                elif isinstance(value, datetime.date):
                    row[key] = value.isoformat()
        return rows
    except Exception as error:
        raise ValidationError("Use UTF-8 CSV or XLSX with entry_id, date, account_code, debit and credit columns; at most 5,000 rows and 4 MB. Dates must use YYYY-MM-DD and amounts use decimal points.") from error


class BooksFileIntake(models.Model):
    _inherit = "poseidon.books.intake"

    @api.model
    def _history_scope(self, company_ids):
        if not isinstance(company_ids, list) or not 1 <= len(company_ids) <= 10 or any(type(i) is not int or i < 1 for i in company_ids):
            raise ValidationError("Select one to ten authorized companies.")
        if not set(company_ids).issubset(self.env.user.company_ids.ids):
            raise AccessError("Company scope is not authorized.")
        return self.with_context(allowed_company_ids=company_ids)

    @api.model
    def historical_readiness(self, company_ids):
        scoped = self._history_scope(company_ids)
        scoped.check_access("read")
        result = []
        for company in scoped.env["res.company"].browse(company_ids):
            domain = [("company_id", "=", company.id)]
            counts = {state: count for state, count in scoped._read_group(domain, ["state"], ["__count"])}
            decisions = scoped.env["poseidon.mapping.decision"]
            mapping_counts = {state: count for state, count in decisions._read_group(domain, ["state"], ["__count"])} if decisions.has_access("read") else None
            connections = scoped.env["qbo.company.mapping"]
            mappings = connections.search(domain) if connections.has_access("read") else None
            result.append({"company_id": company.id, "company_name": company.name,
                           "intake": {state: counts.get(state, 0) for state in ("received", "blocked", "ready", "applied")},
                           "mapping": mapping_counts, "qbo_connections": len(mappings) if mappings is not None else None,
                           "qbo_sync_enabled": any(m.sync_enabled for m in mappings) if mappings is not None else None,
                           "coverage": "QBO JournalEntry and reviewed CSV/XLSX journal files; not a complete historical transaction migration."})
        return result

    @api.model
    def prepare_batch(self, company_ids, after_id=0, limit=25):
        scoped = self._history_scope(company_ids)
        if type(after_id) is not int or after_id < 0 or type(limit) is not int or not 1 <= limit <= 25:
            raise ValidationError("Preparation requires a nonnegative cursor and at most 25 records.")
        records = scoped.search([("company_id", "in", company_ids), ("state", "!=", "applied"), ("id", ">", after_id)], order="id", limit=limit)
        for record in records:
            record.action_prepare()
        cursor = records[-1].id if records else after_id
        more = bool(scoped.search_count([("company_id", "in", company_ids), ("state", "!=", "applied"), ("id", ">", cursor)], limit=1))
        return {"prepared": len(records), "next_cursor": cursor if more else 0, "has_more": more,
                "readiness": scoped.historical_readiness(company_ids)}

    @api.model
    def stage_file(self, company_id, filename, encoded, source_platform):
        scoped = self._history_scope([company_id])
        scoped.check_access("create")
        if not isinstance(source_platform, str) or not source_platform.strip() or len(source_platform) > 80:
            raise ValidationError("Identify the source platform and organization, up to 80 characters.")
        rows = journal_rows(filename, encoded)
        groups = {}
        for row in rows:
            identity = str(row.get("entry_id") or "").strip()
            if not identity or len(identity) > 120:
                raise ValidationError("Every journal row needs a stable entry_id, up to 120 characters.")
            groups.setdefault(identity, []).append(row)
        if len(groups) > 500:
            raise ValidationError("Upload at most 500 journal entries per file.")
        # ponytail: serialize file revisions per company; refine to source keys if ingestion throughput requires it.
        scoped.env.cr.execute("SELECT id FROM res_company WHERE id = %s FOR UPDATE", [company_id])
        staged = scoped.browse()
        for identity, lines in groups.items():
            payload = {"platform": source_platform.strip(), "entry_id": identity, "rows": lines}
            key = "file:" + digest([source_platform.strip(), identity])
            revision = digest(payload)
            item = scoped.search([("company_id", "=", company_id), ("source_kind", "=", "file_journal"), ("source_key", "=", key), ("source_hash", "=", revision)], limit=1)
            if not item:
                item = super(BooksIntake, scoped).create({"name": "%s journal %s" % (source_platform.strip(), identity), "company_id": company_id,
                    "source_kind": "file_journal", "source_key": key, "source_hash": revision, "payload_json": payload})
            Inbox = scoped.env["poseidon.books.inbox"]
            if not Inbox.search_count([("intake_id", "=", item.id), ("user_id", "=", self.env.uid)]):
                Inbox.create({"intake_id": item.id})
            staged |= item
        return {"received": len(staged), "company_id": company_id, "state": "received", "readiness": scoped.historical_readiness([company_id])}

    def _proposal(self):
        self.ensure_one()
        if self.source_kind != "file_journal":
            return super()._proposal()
        self.check_access("read")
        env = self.with_company(self.company_id).with_context(allowed_company_ids=[self.company_id.id]).env
        payload = self.payload_json
        blockers, lines, dates = [], [], set()
        for index, row in enumerate(payload["rows"]):
            try:
                date = fields.Date.to_date(row.get("date"))
                if not date:
                    raise ValueError("Date required")
                dates.add(date)
                if any(isinstance(row.get(key), bool) for key in ("debit", "credit")):
                    raise ValueError("Boolean amounts are not journal amounts")
                debit, credit = float(row.get("debit") or 0), float(row.get("credit") or 0)
                if not all(math.isfinite(v) and 0 <= v <= 1e12 and abs(v - self.company_id.currency_id.round(v)) < 1e-8 for v in (debit, credit)) or debit and credit:
                    raise ValueError("Invalid journal side or amount")
                if row.get("currency") and row["currency"] != self.company_id.currency_id.name:
                    blockers.append({"kind": "foreign_currency", "line": index})
                code = str(row.get("account_code") or "").strip()
                if isinstance(row.get("account_code"), float) and row["account_code"].is_integer():
                    code = str(int(row["account_code"]))
                account = env["account.account"].search([("company_ids", "in", self.company_id.id), ("code", "=", code)], limit=2)
                if len(account) != 1:
                    blockers.append({"kind": "account_mapping", "line": index, "source_account_name": code})
                    continue
                if account.qbo_id:
                    decision = env["poseidon.mapping.decision"].search([("company_id", "=", self.company_id.id), ("qbo_id", "=", account.qbo_id), ("state", "=", "confirmed")], limit=1)
                    account = decision.destination_account_id
                    if not account or self.company_id not in account.company_ids:
                        blockers.append({"kind": "account_mapping", "line": index, "source_account_name": code})
                        continue
                lines.append((0, 0, {"account_id": account.id, "name": str(row.get("description") or payload["entry_id"])[:500], "debit": debit, "credit": credit}))
            except (ValueError, TypeError, OverflowError):
                blockers.append({"kind": "invalid_date_or_amount", "line": index})
        if len(dates) != 1:
            blockers.append({"kind": "inconsistent_entry_date"})
        for date in dates:
            for field in ("fiscalyear_lock_date", "tax_lock_date", "hard_lock_date", "sale_lock_date", "purchase_lock_date"):
                if field in self.company_id._fields and self.company_id[field] and date <= self.company_id[field]:
                    blockers.append({"kind": "locked_period", "lock": field})
        journal = env["account.journal"].search([("company_id", "=", self.company_id.id), ("type", "=", "general")], limit=1)
        if not journal:
            blockers.append({"kind": "general_journal"})
        if len(lines) < 2 or not any(line[2]["debit"] or line[2]["credit"] for line in lines) or not self.company_id.currency_id.is_zero(sum(line[2]["debit"] - line[2]["credit"] for line in lines)):
            blockers.append({"kind": "unbalanced_or_incomplete"})
        existing = env[self._name].search([("company_id", "=", self.company_id.id), ("source_kind", "=", "file_journal"), ("source_key", "=", self.source_key), ("state", "=", "applied"), ("id", "!=", self.id)], limit=1)
        if existing:
            blockers.append({"kind": "existing_canonical_entry", "move_id": existing.move_id.id})
        vals = None if blockers else {"move_type": "entry", "company_id": self.company_id.id, "journal_id": journal.id,
            "date": next(iter(dates)).isoformat(), "ref": self.name, "line_ids": lines}
        accounts = env["account.account"].browse([line[2]["account_id"] for line in lines])
        return vals, blockers, digest({"source_hash": self.source_hash, "values": vals, "blockers": blockers,
            "accounts": [(a.id, str(a.write_date)) for a in accounts], "journal": (journal.id, str(journal.write_date)) if journal else None,
            "company_write_date": str(self.company_id.write_date)})
